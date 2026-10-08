from flask import Flask, render_template, request, redirect, url_for, session, send_file
import sqlite3
from docx import Document
import os

app = Flask(__name__)

app.secret_key = "rp_generator_secret"

DATABASE = "rp_generator.db"


def get_db():
    conn = sqlite3.connect(DATABASE)
    conn.row_factory = sqlite3.Row
    return conn


def init_db():

    conn = get_db()

    conn.execute("""
        CREATE TABLE IF NOT EXISTS programs (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            title TEXT,
            specialty TEXT,
            qualification TEXT,
            course TEXT,
            group_name TEXT,
            year TEXT,
            semester TEXT,
            language TEXT,
            program_type TEXT,
            code TEXT,
            hours INTEGER,
            credits INTEGER,
            control TEXT,
            goal TEXT,
            tasks TEXT,
            status TEXT
        )
    """)

    conn.execute("""
        CREATE TABLE IF NOT EXISTS topics (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            program_id INTEGER,
            name TEXT,
            theory INTEGER,
            practice INTEGER
        )
    """)

    conn.commit()
    conn.close()


@app.route("/", methods=["GET", "POST"])
def login():

    if request.method == "POST":

        login = request.form["login"]
        password = request.form["password"]

        if login == "teacher" and password == "1234":

            session["user"] = login

            return redirect(url_for("dashboard"))

        return render_template(
            "login.html",
            error="Неверный логин или пароль"
        )

    return render_template("login.html")


@app.route("/logout")
def logout():

    session.clear()

    return redirect(url_for("login"))


@app.route("/dashboard")
def dashboard():

    if "user" not in session:
        return redirect(url_for("login"))

    conn = get_db()

    programs = conn.execute(
        "SELECT * FROM programs ORDER BY id DESC"
    ).fetchall()

    conn.close()

    return render_template(
        "dashboard.html",
        programs=programs
    )


@app.route("/create", methods=["GET", "POST"])
def create():

    if "user" not in session:
        return redirect(url_for("login"))

    if request.method == "POST":

        title = request.form["title"]
        specialty = request.form["specialty"]
        qualification = request.form["qualification"]
        course = request.form["course"]
        group_name = request.form["group_name"]
        year = request.form["year"]
        semester = request.form["semester"]
        language = request.form["language"]
        program_type = request.form["program_type"]
        code = request.form["code"]

        hours = int(request.form["hours"])
        credits = int(request.form["credits"])

        control = request.form["control"]
        goal = request.form["goal"]
        tasks = request.form["tasks"]

        conn = get_db()

        conn.execute("""
            INSERT INTO programs (
                title,
                specialty,
                qualification,
                course,
                group_name,
                year,
                semester,
                language,
                program_type,
                code,
                hours,
                credits,
                control,
                goal,
                tasks,
                status
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, (
            title,
            specialty,
            qualification,
            course,
            group_name,
            year,
            semester,
            language,
            program_type,
            code,
            hours,
            credits,
            control,
            goal,
            tasks,
            "Черновик"
        ))

        conn.commit()

        program_id = conn.execute(
            "SELECT last_insert_rowid()"
        ).fetchone()[0]

        conn.close()

        return redirect(
            url_for(
                "program",
                program_id=program_id
            )
        )

    return render_template("create.html")


@app.route("/program/<int:program_id>")
def program(program_id):

    if "user" not in session:
        return redirect(url_for("login"))

    conn = get_db()

    program_data = conn.execute(
        "SELECT * FROM programs WHERE id = ?",
        (program_id,)
    ).fetchone()

    topics = conn.execute(
        "SELECT * FROM topics WHERE program_id = ?",
        (program_id,)
    ).fetchall()

    conn.close()

    if program_data is None:
        return "Программа не найдена"

    return render_template(
        "program.html",
        program=program_data,
        topics=topics
    )


@app.route(
    "/program/<int:program_id>/add_topic",
    methods=["POST"]
)
def add_topic(program_id):

    if "user" not in session:
        return redirect(url_for("login"))

    name = request.form["name"]

    theory = int(request.form["theory"])
    practice = int(request.form["practice"])

    conn = get_db()

    conn.execute("""
        INSERT INTO topics (
            program_id,
            name,
            theory,
            practice
        )
        VALUES (?, ?, ?, ?)
    """, (
        program_id,
        name,
        theory,
        practice
    ))

    conn.commit()
    conn.close()

    return redirect(
        url_for(
            "program",
            program_id=program_id
        )
    )


@app.route("/topic/<int:topic_id>/delete")
def delete_topic(topic_id):

    if "user" not in session:
        return redirect(url_for("login"))

    conn = get_db()

    topic = conn.execute(
        "SELECT program_id FROM topics WHERE id = ?",
        (topic_id,)
    ).fetchone()

    if topic:

        program_id = topic["program_id"]

        conn.execute(
            "DELETE FROM topics WHERE id = ?",
            (topic_id,)
        )

        conn.commit()
        conn.close()

        return redirect(
            url_for(
                "program",
                program_id=program_id
            )
        )

    conn.close()

    return redirect(url_for("dashboard"))


@app.route("/program/<int:program_id>/generate")
def generate_topics(program_id):

    if "user" not in session:
        return redirect(url_for("login"))

    conn = get_db()

    program_data = conn.execute(
        "SELECT * FROM programs WHERE id = ?",
        (program_id,)
    ).fetchone()

    if program_data is None:
        conn.close()
        return "Программа не найдена"

    conn.execute(
        "DELETE FROM topics WHERE program_id = ?",
        (program_id,)
    )

    topic_names = [
        "Введение в дисциплину",
        "Основные понятия и определения",
        "Теоретические основы",
        "Основные методы работы",
        "Информационные технологии",
        "Работа с программным обеспечением",
        "Практическое применение",
        "Работа с данными",
        "Решение практических задач",
        "Контроль и проверка знаний",
        "Итоговое повторение"
    ]

    hours = program_data["hours"]

    number_of_topics = max(
        1,
        hours // 2
    )

    for i in range(number_of_topics):

        if i < len(topic_names):
            name = topic_names[i]
        else:
            name = "Практическое занятие №" + str(i + 1)

        conn.execute("""
            INSERT INTO topics (
                program_id,
                name,
                theory,
                practice
            )
            VALUES (?, ?, ?, ?)
        """, (
            program_id,
            name,
            1,
            1
        ))

    conn.commit()
    conn.close()

    return redirect(
        url_for(
            "program",
            program_id=program_id
        )
    )


@app.route("/program/<int:program_id>/send")
def send_program(program_id):

    if "user" not in session:
        return redirect(url_for("login"))

    conn = get_db()

    conn.execute("""
        UPDATE programs
        SET status = ?
        WHERE id = ?
    """, (
        "На проверке",
        program_id
    ))

    conn.commit()
    conn.close()

    return redirect(
        url_for(
            "program",
            program_id=program_id
        )
    )


@app.route("/program/<int:program_id>/word")
def word(program_id):

    if "user" not in session:
        return redirect(url_for("login"))

    conn = get_db()

    program_data = conn.execute(
        "SELECT * FROM programs WHERE id = ?",
        (program_id,)
    ).fetchone()

    topics = conn.execute(
        "SELECT * FROM topics WHERE program_id = ?",
        (program_id,)
    ).fetchall()

    conn.close()

    if program_data is None:
        return "Программа не найдена"

    document = Document()

    document.add_heading(
        "РАБОЧАЯ УЧЕБНАЯ ПРОГРАММА",
        level=1
    )

    document.add_paragraph(
        "Дисциплина: " + program_data["title"]
    )

    document.add_paragraph(
        "Специальность: " + program_data["specialty"]
    )

    document.add_paragraph(
        "Квалификация: " + program_data["qualification"]
    )

    document.add_paragraph(
        "Курс: " + program_data["course"]
    )

    document.add_paragraph(
        "Группа: " + program_data["group_name"]
    )

    document.add_paragraph(
        "Учебный год: " + program_data["year"]
    )

    document.add_paragraph(
        "Семестр: " + program_data["semester"]
    )

    document.add_paragraph(
        "Язык обучения: " + program_data["language"]
    )

    document.add_paragraph(
        "Количество часов: " + str(program_data["hours"])
    )

    document.add_paragraph(
        "Количество кредитов: " + str(program_data["credits"])
    )

    document.add_paragraph(
        "Форма контроля: " + program_data["control"]
    )

    document.add_heading(
        "Цель программы",
        level=2
    )

    document.add_paragraph(
        program_data["goal"]
    )

    document.add_heading(
        "Задачи",
        level=2
    )

    document.add_paragraph(
        program_data["tasks"]
    )

    document.add_heading(
        "Тематический план",
        level=2
    )

    table = document.add_table(
        rows=1,
        cols=4
    )

    table.style = "Table Grid"

    table.rows[0].cells[0].text = "№"
    table.rows[0].cells[1].text = "Тема"
    table.rows[0].cells[2].text = "Теория"
    table.rows[0].cells[3].text = "Практика"

    for index, topic in enumerate(topics, start=1):

        row = table.add_row()

        row.cells[0].text = str(index)
        row.cells[1].text = topic["name"]
        row.cells[2].text = str(topic["theory"])
        row.cells[3].text = str(topic["practice"])

    file_name = "RP_" + str(program_id) + ".docx"

    document.save(file_name)

    return send_file(
        file_name,
        as_attachment=True
    )


init_db()


if __name__ == "__main__":

    port = int(
        os.environ.get(
            "PORT",
            5000
        )
    )

    app.run(
        host="0.0.0.0",
        port=port,
        debug=True
    )